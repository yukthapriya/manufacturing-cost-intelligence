"""Build the month-end close package for one period: an Excel workbook plus CSVs.

    python -m exporter.close_package --period 2025-03 --out exports/close

Sheets: Summary (sign-off status and gate), Variance by Plant, Subledger to GL,
Inventory Balances, WIP Roll-forward, Three-Way Exceptions, Inventory Exceptions,
All Open Exceptions. Reads from the warehouse; run after dbt and close.period_close.
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from common.db import query_df
from common.logging_setup import get_logger

log = get_logger("exporter.close_package")

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
MONEY = "#,##0.00;[Red]-#,##0.00"

QUERIES = {
    "Variance by Plant": """
        select plant_id, cost_category, cost_element, count(distinct order_id) as orders,
               round(sum(variance_amount), 2) as variance_amount,
               round(sum(variance_amount) filter (where is_unfavorable), 2) as unfavorable,
               round(sum(variance_amount) filter (where not is_unfavorable), 2) as favorable
        from marts.fct_cost_variance where period_id = %(p)s
        group by 1, 2, 3 order by 1, 2, 3""",
    "Subledger to GL": """
        select plant_id, account, account_name, source_type, subledger_amount, gl_amount, diff, status,
               root_cause, break_date, aging_bucket, is_material
        from recon.recon_subledger_to_gl where period_id = %(p)s
        order by (status = 'matched'), abs(diff) desc, plant_id, account, source_type""",
    "Inventory Balances": """
        select plant_id, account, account_name, subledger_balance, gl_balance, gl_vs_subledger_diff,
               independent_valuation, valuation_vs_subledger_diff, status, is_material
        from recon.recon_inventory_balance where period_id = %(p)s order by plant_id, account""",
    "WIP Roll-forward": """
        select plant_id, beginning_wip, material_issued, labor, overhead, completions, scrap, settlement,
               ending_wip_rollforward, ending_open_orders, rollforward_diff, ties_out, gl_ending_wip, gl_diff
        from recon.recon_wip_rollforward where period_id = %(p)s order by plant_id""",
    "Three-Way Exceptions": """
        select po_id, po_line, supplier_id, plant_id, material_id, po_qty, received_qty, invoiced_qty, po_price,
               invoice_price, qty_diff_pct, price_diff_pct, root_cause, amount_at_risk, age_days, aging_bucket,
               invoice_refs, candidate_invoice_ref
        from recon.recon_three_way_match where period_id = %(p)s and status = 'exception'
        order by abs(amount_at_risk) desc""",
    "Inventory Exceptions": """
        select plant_id, material_id, perpetual_qty, system_qty, counted_qty, system_vs_perpetual_value,
               count_vs_system_value, flag_negative_system_qty, flag_system_break, flag_count_variance
        from marts.fct_inventory_discrepancy
        where period_id = %(p)s and (flag_negative_system_qty or flag_system_break or flag_count_variance)
        order by plant_id, material_id""",
    "All Open Exceptions": """
        select recon_type, plant_id, description, amount, status, root_cause, age_days, aging_bucket, is_critical
        from recon.recon_exceptions where period_id = %(p)s
        order by is_critical desc, abs(amount) desc nulls last""",
}
MONEY_HINTS = (
    "amount",
    "variance",
    "diff",
    "balance",
    "valuation",
    "wip",
    "issued",
    "labor",
    "overhead",
    "completions",
    "scrap",
    "settlement",
    "value",
    "favorable",
    "unfavorable",
    "price",
)


def _write_sheet(wb: Workbook, title: str, df: pd.DataFrame) -> None:
    ws = wb.create_sheet(title[:31])
    ws.append(list(df.columns))
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in df.itertuples(index=False):
        ws.append([None if pd.isna(v) else (float(v) if hasattr(v, "is_finite") else v) for v in row])
    for i, col in enumerate(df.columns, start=1):
        letter = get_column_letter(i)
        width = max(
            [len(str(col))] + [len(f"{v:,.2f}" if isinstance(v, float) else str(v)) for v in df[col].head(200).tolist()]
        )
        ws.column_dimensions[letter].width = min(max(10, width + 2), 45)
        if any(h in col for h in MONEY_HINTS) and not col.endswith("_pct") and col not in ("po_id",):
            for cell in ws[letter][1:]:
                cell.number_format = MONEY
        if col.endswith("_pct"):
            for cell in ws[letter][1:]:
                cell.number_format = "0.00%"
    ws.freeze_panes = "A2"
    if len(df):
        ws.auto_filter.ref = ws.dimensions


def build(period: str, out_dir: Path) -> Path:
    out_dir.mkdir(parents=True, exist_ok=True)
    status = query_df("select * from close.period_status where period_id = %(p)s", {"p": period})
    frames = {name: query_df(sql, {"p": period}) for name, sql in QUERIES.items()}

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = f"Month-End Close Package - {period}"
    ws["A1"].font = Font(bold=True, size=14)
    st = status.iloc[0].to_dict() if len(status) else {}
    gate = st.get("gate_results") or {}
    if isinstance(gate, str):
        gate = json.loads(gate)
    summary = [
        ("Sign-off status", st.get("status", "not run")),
        ("Locked", bool(st.get("locked", False))),
        ("Signed off at", str(st.get("signed_off_at") or "")),
        ("Signed off by", st.get("signed_off_by") or ""),
        ("Material GL breaks (period activity)", gate.get("material_gl_breaks", {}).get("count")),
        ("WIP plants not tying", gate.get("wip_rollforward", {}).get("not_tying")),
        ("Failing dbt tests", gate.get("dbt_tests", {}).get("failed")),
        ("Material balance breaks carried forward", gate.get("material_balance_breaks_carried", {}).get("count")),
        ("Open exceptions", len(frames["All Open Exceptions"])),
        (
            "Critical exceptions",
            int(frames["All Open Exceptions"]["is_critical"].sum()) if len(frames["All Open Exceptions"]) else 0,
        ),
        (
            "Total variance",
            float(frames["Variance by Plant"]["variance_amount"].sum()) if len(frames["Variance by Plant"]) else 0.0,
        ),
    ]
    for i, (k, v) in enumerate(summary, start=3):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws.cell(row=i, column=2, value=v)
    ws.cell(row=13, column=2).number_format = MONEY
    reasons = st.get("blocked_reasons") or []
    row = len(summary) + 5
    ws.cell(row=row, column=1, value="Blocking reasons").font = Font(bold=True)
    for j, r in enumerate(reasons or ["none"], start=row + 1):
        ws.cell(row=j, column=1, value=r)
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 28

    for name, df in frames.items():
        _write_sheet(wb, name, df)
        df.to_csv(out_dir / f"close_{period}_{name.lower().replace(' ', '_').replace('-', '')}.csv", index=False)

    path = out_dir / f"close_package_{period}.xlsx"
    wb.save(path)
    log.info(
        "close package written",
        extra={"period": period, "path": str(path), "sheets": len(frames) + 1, "status": st.get("status")},
    )
    return path


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--period", required=True)
    ap.add_argument("--out", default="exports/close", type=Path)
    args = ap.parse_args(argv)
    build(args.period, args.out)
    return 0


if __name__ == "__main__":
    sys.exit(main())
